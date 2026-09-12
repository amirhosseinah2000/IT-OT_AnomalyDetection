"""Build the five-page Persian technical report for the anomaly platform."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "technical-system-report-fa.docx"

NAVY = "151B3D"
VIOLET = "6D3E9F"
MAROON = "711E3A"
LIGHT = "F4EFF8"


def set_rtl(paragraph, rtl: bool = True) -> None:
    paragraph_format = paragraph.paragraph_format
    paragraph_format.space_after = Pt(3)
    paragraph_format.line_spacing = 1.08
    props = paragraph._p.get_or_add_pPr()
    bidi = props.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        props.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")
    if rtl:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT


def shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    props.append(fill)


def set_cell_text(cell, text: str, bold: bool = False, color: str = NAVY) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    set_rtl(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "B Nazanin"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "B Nazanin")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_title(document: Document, text: str, subtitle: str | None = None) -> None:
    paragraph = document.add_paragraph()
    set_rtl(paragraph)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "B Nazanin"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "B Nazanin")
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor.from_string(VIOLET)
    if subtitle:
        sub = document.add_paragraph()
        set_rtl(sub)
        sub.paragraph_format.space_after = Pt(8)
        run = sub.add_run(subtitle)
        run.font.name = "B Nazanin"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "B Nazanin")
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string(MAROON)


def add_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    set_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.name = "B Nazanin"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "B Nazanin")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def add_body(document: Document, text: str, *, bullet: bool = False) -> None:
    paragraph = document.add_paragraph()
    set_rtl(paragraph)
    if bullet:
        paragraph.paragraph_format.right_indent = Cm(0.35)
        text = "• " + text
    run = paragraph.add_run(text)
    run.font.name = "B Nazanin"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "B Nazanin")
    run.font.size = Pt(9.2)
    run.font.color.rgb = RGBColor.from_string(NAVY)


def add_code(document: Document, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    shade(cell, LIGHT)
    paragraph = cell.paragraphs[0]
    set_rtl(paragraph, rtl=False)
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.space_before = Pt(1)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(7.4)
    run.font.color.rgb = RGBColor.from_string(MAROON)


def add_table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, header in zip(table.rows[0].cells, headers):
        shade(cell, VIOLET)
        set_cell_text(cell, header, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)


def add_page_number(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("صفحه ")
    run.font.name = "B Nazanin"
    run.font.size = Pt(8)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def page_break(document: Document) -> None:
    document.add_page_break()


def main() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)
    add_page_number(section)

    styles = document.styles
    styles["Normal"].font.name = "B Nazanin"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "B Nazanin")
    styles["Normal"].font.size = Pt(9.2)

    # Page 1 — scope and outcomes.
    add_title(
        document,
        "گزارش فنی سامانهٔ تشخیص ناهنجاری ترافیک شبکه",
        "معماری PCAP-first، نگاشت قابل ممیزی، انتخاب پویای ویژگی و مدل دومرحله‌ای CPU-first",
    )
    add_heading(document, "۱. هدف، محدوده و خروجی قابل تحویل")
    add_body(
        document,
        "سامانه برای تحلیل ترافیک شبکه از فایل PCAP/PCAPNG طراحی شده است. هدف Stage 1 تشخیص حمله یا رفتار غیرعادی و هدف Stage 2 تشخیص نوع حمله است. منبع همهٔ ویژگی‌ها PCAP است؛ CSV هرگز ورودی مدل نیست و فقط نقش evidence برای نگاشت و ارزیابی برچسب را دارد.",
    )
    add_body(document, "پروتکل‌های فعال فعلی: DNS، HTTP، Modbus و S7comm. هر پروتکل، capture، profile و مدل خروجی مستقل دارد.")
    add_body(document, "معماری به‌صورت پیش‌فرض CPU-first است و برای WSL، Ubuntu و انتقال به سرور طراحی شده؛ Parquet/JSON منبع حقیقت هستند و ClickHouse فقط mirror اختیاری است.")
    add_heading(document, "دستاوردهای پیاده‌سازی‌شده")
    add_table(
        document,
        ["حوزه", "پیاده‌سازی"],
        [
            ["دیتاست", "استخراج جریانی و فایل‌های جدا برای protocol/capture"],
            ["نگاشت", "رابطهٔ packet↔CSV، cache نسخه‌دار و audit قابل ممیزی"],
            ["ویژگی", "profile immutable، quality report و قرارداد inference"],
            ["مدل", "LSTM-AE + Isolation Forest، سپس Random Forest مشروط"],
            ["خروجی", "Parquet/JSON/ONNX، داشبورد فارسی RTL و artifact catalog"],
        ],
    )
    add_heading(document, "اصل طراحی")
    add_body(document, "هیچ مرحله‌ای به dashboard وابسته نیست. dashboard فقط مصرف‌کنندهٔ artifactهای ذخیره‌شده است؛ بنابراین همان نتایج در API، سرویس خارجی، notebook یا سیستم ClickHouse نیز قابل استفاده‌اند.")
    page_break(document)

    # Page 2 — ingestion and mapping.
    add_title(document, "۲. جریان داده، استخراج و نگاشت برچسب")
    add_heading(document, "ساختار و اجرای دیتاست")
    add_body(document, "برای هر پروتکل، ترافیک benign و attack در مسیرهای جدا پردازش می‌شود. خروجی benign در `features/benign/<protocol>/records.parquet` و خروجی attack به تفکیک protocol/capture ذخیره می‌شود. در حالت `--all-packets` استخراج stream شده و row groupهای Parquet به‌صورت دوره‌ای نوشته می‌شوند تا PCAP چندمیلیونی در RAM جمع نشود.")
    add_code(document, "dataset inspect  →  dataset extract --all-packets  →  dataset map <RUN>")
    add_heading(document, "نگاشت CSV با packet یا flow")
    add_body(document, "CSVهای مرجع ممکن است برچسب flow، بازهٔ زمانی یا event داشته باشند. سامانه ابتدا endpoint، پروتکل، زمان و evidenceهای سازگار را بررسی می‌کند؛ سپس همین رابطه را به packetهای واقعیِ استخراج‌شده متصل می‌سازد. بنابراین CSV یک ادعای label است، اما `packet_uid` و `packet_index` مشخص می‌کنند label دقیقاً به کدام packet PCAP تعلق گرفته است.")
    add_body(document, "تمام packetهای ذخیره‌شده batch-by-batch بررسی می‌شوند؛ نمونهٔ محدود extractor فقط برای evidence اولیه است و جای بررسی کامل packetها را نمی‌گیرد.")
    add_heading(document, "نگاشت یک‌بار، استفادهٔ چندباره")
    add_body(document, "اولین اجرای `dataset map` relation و cache نسخه‌دار را می‌سازد. اگر PCAP، CSV و policy تغییر نکرده باشند، اجراهای بعدی از cache استفاده می‌کنند. آموزش مدل هرگز mapping را تکرار نمی‌کند و تنها `labelled/.../records.parquet` را می‌خواند.")
    add_table(
        document,
        ["artifact", "کاربرد"],
        [
            ["mappings/*.audit.json", "نرخ پذیرش، confidence، offset زمانی و علت رد"],
            ["reports/mapping-audit.parquet", "نمای کلی نگاشت همهٔ protocolها"],
            ["labelled/attack/.../records.parquet", "فیچر PCAP همراه برچسب packet-level برای Stage 2"],
            ["mapping_cache/.../packet-csv-links.parquet", "relation قابل استفاده در سیستم خارجی"],
        ],
    )
    add_body(document, "اگر PCAP تغییر کند باید RUN جدید ساخته شود. اگر فقط CSV یا policy تغییر کرده باشد، `dataset map <RUN> --protocol <name> --remap` کافی است.")
    page_break(document)

    # Page 3 — profiles/preprocessing.
    add_title(document, "۳. انتخاب پویای ویژگی و پیش‌پردازش قابل بازتولید")
    add_heading(document, "پروفایل ویژگی")
    add_body(document, "کاربر برای هر پروتکل می‌تواند profile بسازد. profile شامل نام، نسخه، protocolهای مجاز، فهرست ویژگی‌ها و دلیل انتخاب است. profile immutable است؛ تغییر فهرست ویژگی‌ها نسخهٔ جدید می‌سازد تا آموزش و inference قابل بازتولید باشند.")
    add_body(document, "دو الگوی عملی برای هر پروتکل آماده شده‌اند: `*-special-v1` برای ویژگی‌های کم‌هزینه و معنادار همان پروتکل، و `*-all-v1` برای تمام ویژگی‌های catalogue مجاز همان پروتکل.")
    add_heading(document, "گزارش کیفیت و قواعد حذف")
    add_body(document, "فرمان `select analyze` برای هر protocol گزارش می‌سازد: coverage، missing ratio، unique count، variance، IQR، zero ratio، هزینهٔ تقریبی و وضعیت استخراج. ویژگی‌های `constant`، `not_observed` و `not_implemented` برای مدل نهایی مناسب نیستند.")
    add_code(document, "select analyze <records.parquet> --output <profile-quality.parquet>")
    add_heading(document, "pipeline پیش‌پردازش")
    add_body(document, "در آموزش، profile به ماتریس عددی تبدیل می‌شود: ستون‌های غایب و ثابت ثبت و حذف می‌شوند، مقادیر عددی پاک‌سازی/clip و impute می‌شوند، دسته‌های پرکاردینالیتی به واژگان کنترل‌شده و `__OTHER__` تبدیل می‌شوند و ترتیب دقیق ستون‌ها حفظ می‌شود.")
    add_body(document, "خروجی `preprocessing.manifest.json` و pipeline ذخیره‌شده، selected input features، ویژگی‌های حذف‌شده، vocabulary دسته‌ای و تعداد ستون‌های transformed را ثبت می‌کنند. همین قرارداد در inference استفاده می‌شود؛ بنابراین تغییر دستی ترتیب یا featureها مجاز نیست.")
    add_table(
        document,
        ["لایه", "مسئولیت"],
        [
            ["Feature extractor", "تولید ویژگی از PCAP"],
            ["Quality report", "بررسی coverage و variation برای انتخاب کاربر"],
            ["Feature profile", "انتخاب نسخه‌دار featureهای خام"],
            ["Preprocessing manifest", "تبدیل دقیق و قرارداد deploy"],
        ],
    )
    page_break(document)

    # Page 4 — models and evaluation.
    add_title(document, "۴. طراحی مدل دومرحله‌ای و ارزیابی")
    add_heading(document, "Stage 1: تشخیص ناهنجاری")
    add_body(document, "Stage 1 با benign PCAP آموزش می‌بیند و از LSTM Autoencoder برای reconstruction error و Isolation Forest برای isolation score استفاده می‌کند. دو score نسبت به thresholdهای calibration نرمال می‌شوند و تصمیم نهایی با threshold کالیبره‌شدهٔ مشترک گرفته می‌شود. هدف پیش‌فرض false-positive-rate برابر ۰٫۵٪ است.")
    add_body(document, "دادهٔ normal به train، calibration و evaluation دست‌نخورده تقسیم می‌شود. calibration برای threshold است و evaluation برای گزارش FPR واقعی؛ بنابراین معیار گزارش‌شده optimistic نیست.")
    add_heading(document, "Stage 2: تشخیص نوع حمله")
    add_body(document, "Random Forest فقط روی packetهای attack با `mapping_accepted=true` آموزش می‌بیند و علاوه بر ویژگی‌های انتخاب‌شده، می‌تواند scoreهای Stage 1 را بگیرد. `balanced_subsample` برای کلاس‌های نامتوازن فعال است.")
    add_body(document, "پیش‌نیازهای پیش‌فرض Stage 2: حداقل ۲۰ packet، حداقل دو attack type و حداقل چهار نمونه از هر کلاس. نبود این شرایط خطا نیست؛ RF با status=`skipped` ثبت می‌شود و Stage 1 به‌تنهایی deploy می‌شود.")
    add_heading(document, "آزمون hold-out و کنترل نشت داده")
    add_body(document, "برای Stage 2 ابتدا تلاش می‌شود یک یا چند capture کامل برای test نگه داشته شود. اگر این جداسازی باعث حذف یک class از train شود، fallback stratified در سطح packet استفاده و در metrics ثبت می‌شود. فایل `stage2/held-out-split.parquet` مرجع قطعی عضویت train/test است.")
    add_table(
        document,
        ["لایه", "معیارهای کلیدی"],
        [
            ["Stage 1", "FPR، precision، recall، F1، ROC-AUC، PR-AUC و threshold"],
            ["Stage 2", "weighted F1، accuracy، confusion matrix و class metrics"],
            ["End-to-end", "عبور از Stage 1 و درست‌بودن هم‌زمان نوع حمله"],
            ["منابع", "زمان آموزش، RAM، CPU و inference latency"],
        ],
    )
    add_body(document, "در حالت Stage 1-only، anomalyهای عبورکرده از gate با `stage1_anomaly_untyped` ذخیره می‌شوند؛ داشبورد فقط نمودارها و معیارهای واقعی Stage 1 را نمایش می‌دهد.")
    page_break(document)

    # Page 5 — operations, artifacts, and next steps.
    add_title(document, "۵. خروجی‌ها، عملیات، داشبورد و گام‌های بهره‌برداری")
    add_heading(document, "ساختار artifact هر مدل")
    add_code(
        document,
        "models/<protocol>-<profile>/\n"
        "  stage1/metrics.json | scores.parquet | feature-evidence.parquet\n"
        "  stage2/readiness.json | metrics.json | held-out-split.parquet\n"
        "  pipeline/end-to-end-evaluation.parquet | inference-latency.parquet\n"
        "  model-contract.json | onnx/"
    )
    add_body(document, "تمام نمودارهای dashboard بر اساس همین فایل‌ها ساخته می‌شوند. `catalog/artifacts.parquet` فهرست مرکزی مسیرها و artifactهای قابل استفادهٔ بیرون از dashboard است. مدل اصلی و نسخهٔ ONNX همراه contract ذخیره می‌شوند.")
    add_heading(document, "داشبورد و لاگ عملیاتی")
    add_body(document, "داشبورد فارسی و راست‌چین، وضعیت داده/نگاشت، quality features، profileها، آموزش پس‌زمینه، منابع CPU/RAM، خروجی Stage 1، اهمیت featureها، نتایج RF و ارزیابی end-to-end را نمایش می‌دهد. console نیز رویدادهای انگلیسی `DATASET_PROGRESS`، `STREAM_*` و training progress را ثبت می‌کند.")
    add_code(document, "uv run anomaly --config config/production-streaming.yaml dashboard")
    add_heading(document, "روال بهره‌برداری پیشنهادی")
    add_body(document, "۱) extract کامل و map یک‌بار؛ ۲) بازبینی mapping audit؛ ۳) ساخت quality report و profile؛ ۴) آموزش مدل‌های special/all برای هر protocol؛ ۵) مقایسهٔ FPR، Recall، زمان و Stage 2 weighted F1؛ ۶) انتخاب profile نهایی و deploy تنها با model-contract همان مدل.", bullet=True)
    add_body(document, "برای دیتاست حجیم، کل corpus روی دیسک استخراج و نگاشت می‌شود؛ `models.max_source_rows` تنها سقف نمونهٔ یکنواخت fit است تا RAM کنترل شود. افزایش آن باید مرحله‌ای و پس از مشاهدهٔ dashboard/log انجام شود؛ مقدار null برای دیتاست چندمیلیونی مسیر امنی نیست.")
    add_heading(document, "وضعیت اعتبارسنجی و محدودیت‌های باقی‌مانده")
    add_body(document, "تست‌های mapping، cache و protocol-specific quality با ۱۴ تست موفق اجرا شده‌اند و lint/compile بخش‌های آموزش و dashboard بدون خطا است. معیارهای نهایی دقت باید پس از اجرای کامل corpus واقعی و بررسی hold-out هر protocol گزارش شوند؛ این گزارش ادعای accuracy نهایی بدون آن اجرا نمی‌کند.")
    add_body(document, "راهنماهای تکمیلی پروژه شامل انتخاب پویای ویژگی، پردازش هر feature، استفاده از خروجی‌ها در سیستم خارجی، معنی نمودارها و معیارهای ارزیابی در پوشهٔ `guides/` هستند.")

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
