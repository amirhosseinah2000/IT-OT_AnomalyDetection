"""Generate two formal ten-page Persian technical reports for this platform.

The reports intentionally use local, code-drawn diagrams.  They document the
implemented contracts and do not depend on dashboard screenshots or external
services, so the documents remain portable and reproducible offline.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"

NAVY = "18213F"
VIOLET = "693B99"
MAROON = "6D203B"
BLUE = "245CA6"
PALE = "F2EDF7"
WHITE = "FFFFFF"
GRAY = "5A6172"
FONT_FA = "B Nazanin"


def _pil_font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    """Return a practical local font for English technical diagram labels."""
    candidates = [
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
        "C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf",
    ]
    for candidate in candidates:
        if Path(candidate).exists():
            return ImageFont.truetype(candidate, size)
    return ImageFont.load_default()


def _rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    subtitle: str,
    fill: str = "F6F1FB",
    outline: str = VIOLET,
) -> None:
    draw.rounded_rectangle(box, radius=22, fill=f"#{fill}", outline=f"#{outline}", width=4)
    x1, y1, x2, _y2 = box
    heading = _pil_font(28, bold=True)
    body = _pil_font(18)
    draw.text((x1 + 18, y1 + 16), title, font=heading, fill=f"#{NAVY}")
    draw.multiline_text((x1 + 18, y1 + 55), subtitle, font=body, fill=f"#{GRAY}", spacing=4)


def _arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    color: str = BLUE,
) -> None:
    draw.line([start, end], fill=f"#{color}", width=5)
    x, y = end
    if abs(end[0] - start[0]) >= abs(end[1] - start[1]):
        sign = 1 if end[0] >= start[0] else -1
        points = [(x, y), (x - sign * 18, y - 11), (x - sign * 18, y + 11)]
    else:
        sign = 1 if end[1] >= start[1] else -1
        points = [(x, y), (x - 11, y - sign * 18), (x + 11, y - sign * 18)]
    draw.polygon(points, fill=f"#{color}")


def _diagram_canvas(title: str) -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", (1800, 940), "#FFFFFF")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, 1800, 94), fill=f"#{NAVY}")
    draw.text((48, 26), title, font=_pil_font(36, bold=True), fill="#FFFFFF")
    return image, draw


def _save(image: Image.Image, name: str) -> Path:
    ASSETS.mkdir(parents=True, exist_ok=True)
    path = ASSETS / name
    image.save(path, "PNG", optimize=True)
    return path


def build_diagrams() -> dict[str, Path]:
    """Create the local diagrams inserted into the reports."""
    diagrams: dict[str, Path] = {}

    image, draw = _diagram_canvas("Two-Stage Detection Architecture")
    boxes = [
        (70, 310, 330, 490, "PCAP", "packet / flow\nfeatures", "EAF4FF", BLUE),
        (430, 220, 740, 390, "STAGE 1", "LSTM Autoencoder\nIsolation Forest", "F6F1FB", VIOLET),
        (430, 550, 740, 720, "KNOWN-ATTACK GATE", "optional HGB\ntrusted labels only", "FFF1F5", MAROON),
        (870, 310, 1160, 490, "DECISION", "dynamic threshold\nanomaly or normal", "F6F1FB", VIOLET),
        (1290, 220, 1600, 390, "STAGE 2", "Random Forest\nattack type", "EAF4FF", BLUE),
        (1290, 550, 1600, 720, "ARTIFACTS", "scores, evidence\ncontract, ONNX", "FFF1F5", MAROON),
    ]
    for x1, y1, x2, y2, title, subtitle, fill, outline in boxes:
        _rounded_box(draw, (x1, y1, x2, y2), title, subtitle, fill, outline)
    _arrow(draw, (330, 400), (430, 305))
    _arrow(draw, (330, 430), (430, 635), MAROON)
    _arrow(draw, (740, 305), (870, 370))
    _arrow(draw, (740, 635), (870, 430), MAROON)
    _arrow(draw, (1160, 370), (1290, 305))
    _arrow(draw, (1160, 440), (1290, 635), MAROON)
    draw.text((830, 770), "Novel traffic remains covered by Stage 1 even when Stage 2 is unavailable.", font=_pil_font(20), fill=f"#{GRAY}")
    diagrams["two_stage"] = _save(image, "two-stage-architecture.png")

    image, draw = _diagram_canvas("LSTM Autoencoder: Boundary-Safe Reconstruction")
    _rounded_box(draw, (70, 290, 320, 520), "NORMAL SEQUENCES", "prepared features\nwithin one capture", "EAF4FF", BLUE)
    _rounded_box(draw, (450, 220, 760, 420), "BI-LSTM ENCODER", "contextual sequence\nrepresentation", "F6F1FB", VIOLET)
    _rounded_box(draw, (450, 550, 760, 750), "LATENT SPACE", "compact normal\nbehaviour state", "FFF1F5", MAROON)
    _rounded_box(draw, (900, 220, 1210, 420), "LSTM DECODER", "reconstructs each\nfeature sequence", "F6F1FB", VIOLET)
    _rounded_box(draw, (1340, 290, 1660, 520), "RECONSTRUCTION MSE", "packet score\nnormalised threshold", "EAF4FF", BLUE)
    _arrow(draw, (320, 400), (450, 320))
    _arrow(draw, (605, 420), (605, 550), MAROON)
    _arrow(draw, (760, 650), (900, 370), MAROON)
    _arrow(draw, (1210, 320), (1340, 400))
    draw.text((75, 830), "A sequence never crosses a PCAP capture boundary; short capture tails are padded deterministically.", font=_pil_font(22), fill=f"#{GRAY}")
    diagrams["lstm"] = _save(image, "lstm-autoencoder-architecture.png")

    image, draw = _diagram_canvas("Stage 1 Validation and Threshold Policy")
    _rounded_box(draw, (60, 270, 340, 500), "NORMAL CAPTURES", "capture-disjoint\ntrain / calibration / test", "EAF4FF", BLUE)
    _rounded_box(draw, (470, 270, 760, 500), "CALIBRATION", "target FPR\nquantile + MAD guardrail", "F6F1FB", VIOLET)
    _rounded_box(draw, (890, 190, 1190, 390), "LABELLED ATTACK", "separate capture\ntrain / calibration / test", "FFF1F5", MAROON)
    _rounded_box(draw, (890, 550, 1190, 750), "UNLABELLED MIXED", "score and retain\nno accuracy claim", "F7F7FA", GRAY)
    _rounded_box(draw, (1330, 270, 1670, 500), "VALIDATION SUMMARY", "FPR, recall, PR-AUC\nscope stated explicitly", "EAF4FF", BLUE)
    _arrow(draw, (340, 385), (470, 385))
    _arrow(draw, (760, 385), (1330, 385))
    _arrow(draw, (1190, 290), (1330, 350), MAROON)
    _arrow(draw, (1190, 650), (1330, 430), GRAY)
    draw.text((90, 840), "Only trusted CSV-mapped packets contribute to accuracy metrics. Unknown labels are observations, never fabricated negatives.", font=_pil_font(20), fill=f"#{GRAY}")
    diagrams["validation"] = _save(image, "stage1-validation-architecture.png")

    image, draw = _diagram_canvas("Complete Module Architecture")
    nodes = [
        (55, 250, 285, 470, "DATA", "PCAP / PCAPNG\nCSV labels", "EAF4FF", BLUE),
        (360, 250, 610, 470, "DISCOVERY", "inventory\nprotocol folders", "F6F1FB", VIOLET),
        (685, 250, 935, 470, "EXTRACTION", "causal features\nParquet batches", "EAF4FF", BLUE),
        (1010, 250, 1260, 470, "MAPPING", "packet↔CSV audit\ncache", "FFF1F5", MAROON),
        (1335, 250, 1585, 470, "MODELS", "profile + preprocess\ntwo-stage pipeline", "F6F1FB", VIOLET),
        (1335, 620, 1585, 800, "CONSUMERS", "dashboard / CLI\nAPI / ClickHouse", "EAF4FF", BLUE),
    ]
    for x1, y1, x2, y2, title, subtitle, fill, outline in nodes:
        _rounded_box(draw, (x1, y1, x2, y2), title, subtitle, fill, outline)
    for start_x, end_x in [(285, 360), (610, 685), (935, 1010), (1260, 1335)]:
        _arrow(draw, (start_x, 360), (end_x, 360))
    _arrow(draw, (1460, 470), (1460, 620))
    draw.rounded_rectangle((360, 590, 1260, 805), radius=28, fill="#F7F7FA", outline=f"#{GRAY}", width=3)
    draw.text((405, 635), "PORTABLE ARTIFACT LAYER", font=_pil_font(30, bold=True), fill=f"#{NAVY}")
    draw.text((405, 690), "catalog  |  Parquet  |  JSON  |  model contract  |  native models  |  ONNX", font=_pil_font(23), fill=f"#{GRAY}")
    diagrams["module"] = _save(image, "complete-module-architecture.png")

    image, draw = _diagram_canvas("Data, Mapping, and Cache Lifecycle")
    _rounded_box(draw, (70, 260, 340, 490), "EXTRACT ONCE", "PCAP -> features\nper protocol/capture", "EAF4FF", BLUE)
    _rounded_box(draw, (470, 260, 760, 490), "MAP ONCE", "endpoint + time\npacket-level labels", "FFF1F5", MAROON)
    _rounded_box(draw, (890, 260, 1180, 490), "AUDIT", "confidence, offset\naccept / reject", "F6F1FB", VIOLET)
    _rounded_box(draw, (1310, 260, 1630, 490), "REUSE", "labelled Parquet\ncache relation", "EAF4FF", BLUE)
    _arrow(draw, (340, 375), (470, 375))
    _arrow(draw, (760, 375), (890, 375), MAROON)
    _arrow(draw, (1180, 375), (1310, 375))
    draw.rounded_rectangle((470, 620, 1180, 770), radius=24, fill="#F7F7FA", outline=f"#{GRAY}", width=3)
    draw.text((520, 660), "TRAINING READS SAVED LABELLED FEATURES — IT DOES NOT REMAP PCAP/CSV", font=_pil_font(22, bold=True), fill=f"#{MAROON}")
    diagrams["mapping"] = _save(image, "mapping-cache-lifecycle.png")

    image, draw = _diagram_canvas("Feature Profile and Inference Contract")
    _rounded_box(draw, (75, 270, 370, 500), "QUALITY REPORT", "coverage / variance\nfeature applicability", "EAF4FF", BLUE)
    _rounded_box(draw, (485, 270, 780, 500), "PROFILE JSON", "selected raw fields\ncontent-hash version", "F6F1FB", VIOLET)
    _rounded_box(draw, (895, 270, 1190, 500), "PREPROCESSOR", "impute / encode\nrobust scale", "FFF1F5", MAROON)
    _rounded_box(draw, (1305, 270, 1650, 500), "MODEL CONTRACT", "column order + models\nthresholds + ONNX", "EAF4FF", BLUE)
    _arrow(draw, (370, 385), (485, 385))
    _arrow(draw, (780, 385), (895, 385), MAROON)
    _arrow(draw, (1190, 385), (1305, 385))
    draw.text((160, 740), "Inference uses the saved contract exactly; a dashboard selection cannot change a deployed feature set.", font=_pil_font(24), fill=f"#{GRAY}")
    diagrams["contract"] = _save(image, "feature-contract-architecture.png")

    image, draw = _diagram_canvas("Artifact Consumers and Operational Monitoring")
    _rounded_box(draw, (90, 250, 420, 480), "ARTIFACT STORE", "Parquet / JSON\nmodel contract / ONNX", "F6F1FB", VIOLET)
    _rounded_box(draw, (555, 170, 870, 360), "DASHBOARD", "RTL visual analysis\nresource monitor", "EAF4FF", BLUE)
    _rounded_box(draw, (555, 550, 870, 740), "CLI", "English progress logs\nrepeatable commands", "EAF4FF", BLUE)
    _rounded_box(draw, (1030, 170, 1360, 360), "EXTERNAL API", "AnomalyService\ncontract-based score", "FFF1F5", MAROON)
    _rounded_box(draw, (1030, 550, 1360, 740), "CLICKHOUSE", "optional mirror\nnever the source of truth", "FFF1F5", MAROON)
    _rounded_box(draw, (1490, 360, 1740, 550), "OPERATIONS", "alerts, reports\nmodel review", "F6F1FB", VIOLET)
    for end in [(555, 265), (555, 645), (1030, 265), (1030, 645)]:
        _arrow(draw, (420, 365), end)
    _arrow(draw, (1360, 265), (1490, 410), MAROON)
    _arrow(draw, (1360, 645), (1490, 500), MAROON)
    diagrams["operations"] = _save(image, "artifact-consumers-monitoring.png")

    return diagrams


def _set_rtl(paragraph, rtl: bool = True) -> None:
    """Make the paragraph explicitly Persian/RTL in Word.

    Word can otherwise place mixed Persian/English runs left-to-right even
    when the paragraph alignment alone appears correct.
    """
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.08
    paragraph.paragraph_format.keep_together = True
    props = paragraph._p.get_or_add_pPr()
    bidi = props.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        props.append(bidi)
    bidi.set(qn("w:val"), "1" if rtl else "0")


def _set_style_rtl(style) -> None:
    """Make the document's default body style RTL as well."""
    props = style._element.get_or_add_pPr()
    bidi = props.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        props.append(bidi)
    bidi.set(qn("w:val"), "1")


def _set_run_rtl(run, rtl: bool = True) -> None:
    """Apply Complex Script direction and Persian language to a text run."""
    props = run._element.get_or_add_rPr()
    rtl_node = props.find(qn("w:rtl"))
    if rtl_node is None:
        rtl_node = OxmlElement("w:rtl")
        props.append(rtl_node)
    rtl_node.set(qn("w:val"), "1" if rtl else "0")
    language = props.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        props.append(language)
    language.set(qn("w:bidi"), "fa-IR")


def _set_table_rtl(table) -> None:
    """Reverse visual table order and prevent split rows in Word."""
    props = table._tbl.tblPr
    bidi_visual = props.find(qn("w:bidiVisual"))
    if bidi_visual is None:
        bidi_visual = OxmlElement("w:bidiVisual")
        props.append(bidi_visual)
    bidi_visual.set(qn("w:val"), "1")
    for row in table.rows:
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))


def _font(run, size: float = 10, bold: bool = False, color: str = NAVY, name: str = FONT_FA) -> None:
    run.font.name = name
    props = run._element.get_or_add_rPr()
    fonts = props.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        props.insert(0, fonts)
    for font_kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{font_kind}"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    _set_run_rtl(run)


def _shade(cell, color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    props.append(shading)


def _cell(cell, value: str, *, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    _set_rtl(paragraph)
    run = paragraph.add_run(value)
    _font(run, 8.7, bold=header, color=WHITE if header else NAVY)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if header:
        _shade(cell, VIOLET)


def _setup(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.35)
    section.right_margin = Cm(1.35)

    styles = document.styles
    styles["Normal"].font.name = FONT_FA
    normal_props = styles["Normal"]._element.get_or_add_rPr()
    normal_fonts = normal_props.rFonts
    if normal_fonts is None:
        normal_fonts = OxmlElement("w:rFonts")
        normal_props.insert(0, normal_fonts)
    for font_kind in ("ascii", "hAnsi", "eastAsia", "cs"):
        normal_fonts.set(qn(f"w:{font_kind}"), FONT_FA)
    _set_style_rtl(styles["Normal"])
    styles["Normal"].font.size = Pt(10)
    document.core_properties.title = title
    document.core_properties.subject = "Network Anomaly Detection Platform"
    document.core_properties.author = "Network Analytics Team"

    header = section.header.paragraphs[0]
    _set_rtl(header)
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("سامانهٔ تشخیص ناهنجاری ترافیک شبکه  |  گزارش فنی")
    _font(run, 8, color=VIOLET)

    footer = section.footer.paragraphs[0]
    _set_rtl(footer)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("محرمانه — استفادهٔ فنی داخلی | صفحه ")
    _font(run, 8, color=GRAY)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def _title(document: Document, text: str, subtitle: str | None = None) -> None:
    paragraph = document.add_paragraph()
    _set_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(4)
    run = paragraph.add_run(text)
    _font(run, 17, bold=True, color=VIOLET)
    if subtitle:
        sub = document.add_paragraph()
        _set_rtl(sub)
        run = sub.add_run(subtitle)
        _font(run, 9.5, color=MAROON)


def _heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    _set_rtl(paragraph)
    paragraph.paragraph_format.space_before = Pt(5)
    run = paragraph.add_run(text)
    _font(run, 11.5, bold=True, color=NAVY)


def _body(document: Document, text: str, *, bullet: bool = False) -> None:
    paragraph = document.add_paragraph()
    _set_rtl(paragraph)
    if bullet:
        paragraph.paragraph_format.right_indent = Cm(0.4)
        text = "• " + text
    run = paragraph.add_run(text)
    _font(run, 9.6)


def _table(document: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    _set_table_rtl(table)
    for cell, value in zip(table.rows[0].cells, headers, strict=True):
        _cell(cell, value, header=True)
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row, strict=True):
            _cell(cell, value)


def _diagram(document: Document, path: Path, caption: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Cm(15.6))
    cap = document.add_paragraph()
    _set_rtl(cap)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = cap.add_run(caption)
    _font(run, 8.5, color=GRAY)


def _page_break(document: Document, page: int) -> None:
    if page < 10:
        document.add_page_break()


def _cover(document: Document, title: str, subtitle: str, report_id: str) -> None:
    for _ in range(5):
        document.add_paragraph()
    title_p = document.add_paragraph()
    _set_rtl(title_p)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(title)
    _font(run, 22, bold=True, color=VIOLET)
    sub = document.add_paragraph()
    _set_rtl(sub)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run(subtitle)
    _font(run, 12, color=MAROON)
    document.add_paragraph()
    table = document.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    _set_table_rtl(table)
    values = [
        ("شناسهٔ سند", report_id),
        ("نسخه", "۱٫۰ — بر پایهٔ پیاده‌سازی فعلی"),
        ("دامنه", "سامانهٔ تشخیص ناهنجاری شبکه مبتنی بر PCAP"),
        ("مخاطب", "تیم فنی، ارزیاب معماری و بهره‌بردار سامانه"),
    ]
    for row, (right, left) in zip(table.rows, values, strict=True):
        _cell(row.cells[0], right, header=True)
        _cell(row.cells[1], left)
    _body(document, "این گزارش به‌صورت آفلاین و از روی قراردادها و کدهای فعلی پروژه تولید شده است. نمودارهای داخل سند، مسیرهای پیاده‌سازی‌شده را نشان می‌دهند و صرفاً تصویر مفهومی نیستند.")


def build_model_report(diagrams: dict[str, Path]) -> Path:
    document = Document()
    _setup(document, "مستند طراحی مدل‌های تشخیص ناهنجاری")
    output = ROOT / "model-design-technical-report-fa.docx"

    _cover(
        document,
        "مستند طراحی و معماری مدل‌های تشخیص ناهنجاری",
        "LSTM-AE، Isolation Forest، گیت حملهٔ شناخته‌شده و Random Forest دومرحله‌ای",
        "NAD-MODEL-ARCH-FA-001",
    )
    _page_break(document, 1)

    _title(document, "۱. مسئله، اهداف و اصول طراحی مدل")
    _heading(document, "صورت مسئله")
    _body(document, "سامانه باید از featureهای استخراج‌شده از PCAP ابتدا تشخیص دهد آیا یک پکت یا رفتار وابسته به flow ناهنجار است و سپس، فقط در صورت عبور از Stage 1، نوع حمله را تشخیص دهد. دادهٔ attack-side ذاتاً mixed است؛ بنابراین وجود PCAP حمله به معنی حمله‌بودن همهٔ پکت‌ها نیست.")
    _heading(document, "اهداف قابل سنجش")
    for item in [
        "حفظ FPR پایین روی capture نرمالِ دیده‌نشده، بدون تنظیم threshold با test data.",
        "افزایش recall حمله‌های دارای label معتبر، بدون برچسب‌گذاری جعلی پکت‌های unknown.",
        "عملکرد CPU-first، ورودی پویا بر اساس profile و خروجی قابل حمل در native و ONNX.",
        "ارائهٔ evidence برای هر alert، نه صرفاً یک برچسب anomaly.",
    ]:
        _body(document, item, bullet=True)
    _table(document, ["تصمیم", "دلیل فنی"], [["معماری دومرحله‌ای", "کاهش هزینهٔ classifier چندکلاسه روی همهٔ ترافیک و جداسازی کشف از attribution."], ["PCAP-first", "جلوگیری از نشت ستون CSV به featureهای مدل و حفظ قابلیت deployment روی ترافیک خام."], ["profile immutable", "یکسان‌ماندن feature order و transformation میان train و inference."]])
    _page_break(document, 2)

    _title(document, "۲. نمای معماری مدل دومرحله‌ای")
    _body(document, "ورودی Stage 1 ماتریس transform‌شده از profile فریز‌شده است. این مرحله دو signal مکمل می‌سازد: reconstruction error و isolation score. در دادهٔ labelدار کافی، گیت سبکِ supervised برای حمله‌های شناخته‌شده نیز اضافه می‌شود. Stage 2 فقط برای anomalyهای عبورکرده اجرا می‌شود.")
    _diagram(document, diagrams["two_stage"], "شکل ۱ — معماری واقعی مدل و artifactهای خروجی")
    _table(document, ["بخش", "ورودی", "خروجی"], [["Stage 1", "feature matrix و capture group", "score نرمال‌شده، vote، anomaly flag"], ["گیت اختیاری", "normal train + attack labelدار از capture جدا", "known-attack probability"], ["Stage 2", "anomalyهای عبورکرده", "attack type و probability"], ["Evidence", "feature matrix + baseline normal", "رتبهٔ featureهای دور از رفتار normal"]])
    _page_break(document, 3)

    _title(document, "۳. مدل LSTM Autoencoder")
    _heading(document, "چرا LSTM-AE انتخاب شده است؟")
    _body(document, "رفتار شبکه فقط مقدار مستقل یک feature نیست؛ نرخ، فاصلهٔ زمانی، طول پکت، تکرار درخواست و رفتار flow در یک توالی معنا پیدا می‌کنند. LSTM-AE روی توالی normal آموزش می‌بیند و وقتی الگوی جدید قابل reconstruction نباشد، MSE بالاتری تولید می‌کند.")
    _diagram(document, diagrams["lstm"], "شکل ۲ — encoder/latent/decoder و محاسبهٔ reconstruction MSE")
    _table(document, ["جزء", "پیاده‌سازی", "دلیل"], [["Encoder", "Bi-LSTM + LayerNorm", "استفاده از context دوطرفه در window برای نمایش پایدار normal."], ["Latent", "linear + GELU + LayerNorm", "فشرده‌سازی سبک و جلوگیری از latent بسیار بزرگ."], ["Decoder", "LSTM و لایهٔ خروجی feature", "بازسازی توالی در همان فضای transform‌شده."], ["Dynamic architecture", "اندازه براساس تعداد feature", "profile کوچک مدل سبک‌تر و profile بزرگ بدون hard-code."], ["Capture boundary", "sequence_groups", "جلوگیری از ساخت sequence مصنوعی میان دو PCAP نامرتبط."]])
    _page_break(document, 4)

    _title(document, "۴. Isolation Forest و ترکیب سیگنال‌ها")
    _heading(document, "چرا Isolation Forest؟")
    _body(document, "Isolation Forest بدون نیاز به label، نمونه‌های دور از تراکم normal را با مسیرهای کوتاه‌تر در درخت‌ها جدا می‌کند. نسبت به Autoencoder یک failure mode متفاوت دارد: ممکن است پکت با MSE کم هنوز در فضای featureهای normal به‌آسانی isolate شود.")
    _body(document, "مدل روی همان prepared normal matrix fit می‌شود. score آن معکوس می‌گردد تا در تمام لایه‌ها مقدار بزرگ‌تر به معنی ناهنجاری بیشتر باشد. تعداد estimator و worker از config می‌آیند و CPU workerها محدود می‌شوند.")
    _heading(document, "Ensemble و آستانه")
    _body(document, "scoreهای LSTM و Isolation Forest واحد یکسان ندارند. هر score ابتدا بر threshold calibration خودش تقسیم می‌شود. سپس `calibrated_max` یا وزن‌های تعریف‌شده در config یک score بدون‌بُعد می‌سازند. threshold نهایی با distribution نرمال calibration می‌شود، نه با test attack.")
    _table(document, ["حالت", "کاربرد"], [["calibrated_max", "حفظ signal قوی هر detector؛ مناسب وقتی یک detector به یک حمله حساس‌تر است."], ["calibrated_weighted", "ترکیب نرم وزن‌دار و bonus توافق detectorها."], ["both_detectors", "حالت محافظه‌کار برای محیطی که FPR بسیار حساس است."], ["either_detector", "سازگاری با قراردادهای قدیمی؛ حساس‌تر اما ممکن است FPR را بالا ببرد."]])
    _page_break(document, 5)

    _title(document, "۵. گیت حملهٔ شناخته‌شده و کنترل false negative")
    _heading(document, "دلیل افزودن گیت سبک")
    _body(document, "Autoencoder می‌تواند بعضی حمله‌ها را ساده‌تر از traffic normal متنوع بازسازی کند. در این وضعیت پایین‌آوردن threshold فقط false positive را زیاد می‌کند. گیت HistGradientBoosting یک classifier سبک CPU است که فقط از packetهای واقعاً map و label شده استفاده می‌کند و در کنار novelty detector قرار می‌گیرد.")
    _body(document, "گیت تنها وقتی فعال می‌شود که حداقل سه capture دارای attack label معتبر وجود داشته باشد. captureهای مثبت به train، calibration و evaluation تقسیم می‌شوند؛ بنابراین یک capture واحد نمی‌تواند هم یادگیری و هم موفقیت ظاهری مدل را بسازد.")
    _diagram(document, diagrams["validation"], "شکل ۳ — سیاست label، calibration و تفکیک metric از observation")
    _table(document, ["نوع رکورد", "رفتار مدل/metric"], [["normal PCAP", "Stage 1 train، calibration و holdout FPR."], ["trusted mapped attack", "ورود به recall/F1 و در صورت کافی‌بودن captureها آموزش گیت."], ["trusted mapped benign در attack capture", "نمونهٔ منفی معتبر برای بررسی FPR distribution deployment."], ["unknown یا mapping ردشده", "score و ذخیره برای review؛ حذف از TP/TN/FP/FN."]])
    _page_break(document, 6)

    _title(document, "۶. Random Forest مرحلهٔ دوم")
    _heading(document, "نقش و دلیل انتخاب")
    _body(document, "پس از گیت Stage 1، فقط تعداد کمتری candidate به Stage 2 می‌رسند. Random Forest برای tabular featureهای ترکیبی، روابط غیرخطی، مقیاس‌پذیری CPU، پایداری نسبت به scaling و ارائهٔ feature importance انتخاب شده است. این مدل به‌جای anomaly detection، attack-type classification انجام می‌دهد.")
    _body(document, "matrix Stage 2 شامل transformed featureها و در صورت فعال‌بودن، scoreهای LSTM/Isolation Forest است. RF با `balanced_subsample` آموزش می‌بیند تا classهای کم‌تعداد کاملاً نادیده گرفته نشوند.")
    _table(document, ["Readiness", "دلیل"], [["حداقل تعداد رکورد", "جلوگیری از ساخت RF ظاهراً موفق با نمونهٔ بسیار کم."], ["حداقل دو attack type", "تشخیص نوع حملهٔ تک‌کلاسه ارزش classifier ندارد."], ["حداقل نمونه در هر کلاس", "امکان split و metric معنادار."], ["capture-held-out split", "جلوگیری از نشت پکت/flow یک capture در train و test."], ["skip امن", "در نبود label کافی، Stage 1 همچنان deployable است."]])
    _page_break(document, 7)

    _title(document, "۷. ارزیابی، threshold و معیارهای اعتبار")
    _heading(document, "روش تست درست")
    _body(document, "ابتدا FPR روی normal captureهای دست‌نخورده بررسی می‌شود. سپس recall، precision، F1، ROC-AUC و Average Precision فقط روی packetهایی محاسبه می‌شوند که label CSV آن‌ها در mapping پذیرفته شده است. همهٔ scoreها و roleها در `stage1/scores.parquet` قابل audit هستند.")
    _body(document, "حمله‌ای که normal تشخیص داده شود False Negative است؛ normalی که attack اعلام شود False Positive است. این دو هزینهٔ عملیاتی متفاوت دارند و نباید تنها accuracy کلی گزارش شود.")
    _table(document, ["معیار", "سؤال عملیاتی"], [["FPR / Specificity", "چند درصد ترافیک normal باعث alert بی‌مورد می‌شود؟"], ["Recall / FNR", "چند درصد حملهٔ labelدار از Stage 1 عبور نمی‌کند؟"], ["Precision", "از alertها چند مورد واقعاً attack هستند؟"], ["PR-AUC", "با عدم‌توازن شدید، رتبه‌بندی حمله‌ها چقدر مفید است؟"], ["Latency / records/sec", "مدل روی CPU server چه ظرفیت عملیاتی دارد؟"], ["Weighted F1 Stage 2", "نوع حمله در تمام classها چقدر درست تشخیص داده می‌شود؟"]])
    _page_break(document, 8)

    _title(document, "۸. explainability، خروجی و قرارداد inference")
    _heading(document, "دلیل قابل‌توضیح بودن alert")
    _body(document, "خروجی Stage 1 فقط boolean نیست. برای هر packet score LSTM، score Isolation Forest، vote هر detector، threshold، score نهایی، probability گیت و `feature-evidence` ثبت می‌شود. evidence از انحراف robust feature نسبت به median/IQR normal در همان فضای transform‌شده ساخته می‌شود.")
    _body(document, "`model-contract.json` تنها مرجع inference است: مسیر preprocessor، profile، ترتیب ستون‌ها، native modelها، thresholdها و وضعیت Stage 2 را نگه می‌دارد. `score_two_stage` ابتدا contract را می‌خواند و سپس همان preprocessing را اجرا می‌کند؛ از feature order جدید یا dashboard state استفاده نمی‌کند.")
    _table(document, ["artifact", "مصرف بیرونی"], [["model-contract.json", "ورودی سرویس inference یا validator deployment."], ["stage1/scores.parquet", "alert investigation و dashboard/API table."], ["feature-evidence.parquet", "توضیح top featureهای هر alert."], ["pipeline/end-to-end-predictions.parquet", "خروجی نهایی normal/anomaly/type."], ["onnx/*.onnx", "runtimeهای ONNX سازگار، با حفظ contract preprocessing."]])
    _page_break(document, 9)

    _title(document, "۹. منابع، بسته‌بندی و تصمیم‌های عملیاتی")
    _heading(document, "CPU-first و کنترل منابع")
    _body(document, "مدل‌ها برای CPU طراحی شده‌اند. LSTM اندازهٔ hidden/latent را از تعداد feature تعیین می‌کند، `max_train_windows` هزینهٔ train را سقف می‌گذارد، Isolation Forest و RF worker کنترل‌شده دارند و diagnosticهای پرهزینه با estimator محدود اجرا می‌شوند. resource snapshot قبل/بعد train ذخیره می‌شود.")
    _heading(document, "بسته‌بندی و deployment")
    _body(document, "هر model run native artifact و ONNX مستقل دارد. ONNX exporter نتیجهٔ هر جزء را در `onnx-export.json` ثبت می‌کند تا شکست یک converter باعث توقف deployment native نشود. preprocessor joblib و manifest بخشی از مدل هستند، نه فایل فرعی اختیاری.")
    _heading(document, "محدودیت‌ها و تصمیم بعدی")
    for item in [
        "بدون captureهای attack با label معتبر کافی، گیت و Stage 2 عمداً train نمی‌شوند.",
        "اگر recall پایین و FPR پایین باشد، ابتدا mapping coverage و profile باید بررسی شود؛ پایین‌آوردن کورکورانهٔ threshold راه‌حل نیست.",
        "برای Modbus یا protocol کم‌برچسب، افزایش captureهای labelدار و بررسی نمونه‌گیری label-aware اولویت دارد.",
        "نتیجهٔ smoke با ۲۰۰ پکت فقط صحت خروجی را می‌سنجد؛ معیار نهایی از full run و capture-held-out گزارش می‌شود.",
    ]:
        _body(document, item, bullet=True)

    document.save(output)
    return output


def build_module_report(diagrams: dict[str, Path]) -> Path:
    document = Document()
    _setup(document, "مستند معماری کامل ماژول تشخیص ناهنجاری")
    output = ROOT / "module-architecture-technical-report-fa.docx"

    _cover(
        document,
        "مستند معماری کامل ماژول تشخیص ناهنجاری",
        "از کشف داده و نگاشت CSV تا قرارداد مدل، داشبورد، API و استقرار سرور",
        "NAD-MODULE-ARCH-FA-001",
    )
    _page_break(document, 1)

    _title(document, "۱. نمای کلان و اصول معماری ماژول")
    _body(document, "ماژول به مرزهای روشن تقسیم شده است: discovery، extraction، mapping، selection، preprocessing، modelling، storage و presentation. هر مرز artifact خود را می‌نویسد؛ مرحلهٔ بعدی به دادهٔ خام مرحلهٔ قبل وابسته نیست و فایل ذخیره‌شده را مصرف می‌کند.")
    _diagram(document, diagrams["module"], "شکل ۱ — معماری کامل و لایهٔ artifact قابل انتقال")
    _table(document, ["اصل", "نتیجه"], [["Parquet/JSON source of truth", "داشبورد، ClickHouse و API مصرف‌کننده‌اند؛ lock-in ایجاد نمی‌شود."], ["Protocol separation", "اضافه‌کردن protocol یا تحلیل یک protocol بدون قاطی‌شدن داده‌ها."], ["Immutable contracts", "feature/model یک run در inference دقیقاً قابل تکرار است."], ["CPU-first", "سازگاری با WSL، Ubuntu و server بدون GPU."]])
    _page_break(document, 2)

    _title(document, "۲. کشف داده و inventory")
    _heading(document, "ورودی فیزیکی")
    _body(document, "`datasets/inventory.py` پوشه‌های benign/attack/labels را طبق config مرور می‌کند، فایل‌های PCAP/PCAPNG مستقیم را پیدا می‌کند، protocol را از پوشه تعیین می‌کند و برای هر PCAP حمله candidate CSV می‌سازد. allowlist و override برای کنترل دقیق آزمون یا نام‌گذاری غیرهمسان وجود دارد.")
    _body(document, "خروجی `input-inventory.json` باید قبل از هر train بررسی شود: نام capture، protocol، مسیر PCAP، candidateهای CSV و علت حذف احتمالی در آن ثبت می‌شود. این فایل پاسخ می‌دهد آیا دادهٔ جدید واقعاً وارد سامانه شده است یا نه.")
    _table(document, ["موجودیت", "فیلدهای مهم", "مصرف‌کننده"], [["CaptureSource", "path، protocol، split، capture name", "extractor و catalog"], ["AttackPair", "attack PCAP، candidate CSV، protocol", "dataset mapper"], ["allowlist", "نام capture مجاز", "smoke/reprocess هدفمند"], ["override", "pair دستی PCAP/CSV", "دیتاست با نام‌های ناهمسان"]])
    _page_break(document, 3)

    _title(document, "۳. استخراج feature از PCAP")
    _heading(document, "مسئولیت extractor")
    _body(document, "`features/extractor.py` با Scapy پکت را می‌خواند، IP/port/payload/زمان را نرمال می‌کند، protocol واقعی را تشخیص می‌دهد و یک row پایه ایجاد می‌کند. سپس state causal flow، direction، timing، burst و host behavior را به row اضافه می‌کند. هیچ feature به آیندهٔ همان capture نگاه نمی‌کند.")
    _body(document, "برای capture بزرگ، extractor به‌صورت streaming کار می‌کند: row group Parquet نوشته می‌شود، state flow/host سقف دارد، memory کنترل می‌شود و progress انگلیسی در console و JSON ثبت می‌شود. برای smoke sample یکنواخت و قطعی گرفته می‌شود.")
    _table(document, ["لایه", "نمونه output"], [["Packet", "packet_length، payload_size، entropy، TCP flags"], ["Flow", "flow_duration، total packets/bytes، ratio، mean/std"], ["Timing", "inter-arrival، jitter، packet rate، hour/weekday cycles"], ["Host behaviour", "source rate، destination count، entropy، burstiness"], ["Protocol parser", "DNS/HTTP/Modbus/S7comm/SSH fields"]])
    _page_break(document, 4)

    _title(document, "۴. معماری mapping و cache برچسب")
    _body(document, "نگاشت بعد از extraction اجرا می‌شود و همان records ذخیره‌شده را می‌خواند؛ PCAP برای train مجدداً parse نمی‌شود. mapper schemaهای متفاوت CSV IT/OT را normalize می‌کند، endpoint/time سازگار را پیدا می‌کند، offset زمانی محتمل را audit می‌کند و label را در زمان هر packet attach می‌کند.")
    _diagram(document, diagrams["mapping"], "شکل ۲ — lifecycle استخراج، audit نگاشت، cache و استفادهٔ بعدی")
    _body(document, "cache رابطهٔ packet↔CSV را با signature منبع و policy نگه می‌دارد. cache mismatch باعث rebuild همان capture می‌شود، نه این‌که label با ترتیب row فرض شود. mapping پذیرفته‌شده و label معتبر دو مفهوم جدا هستند: یک capture پذیرفته‌شده می‌تواند packetهای unknown داشته باشد.")
    _table(document, ["قانون", "اثر"], [["timestamp + endpoint evidence", "کاهش اتصال اشتباه برچسب CSV به PCAP."], ["mapping audit", "امکان رد رابطهٔ کم‌اعتماد قبل از train."], ["unknown حفظ می‌شود", "عدم تبدیل packet بدون evidence به benign یا attack."], ["cache reuse", "اجرای سریع‌تر و نتیجهٔ قابل تکرار در runهای بعدی."]])
    _page_break(document, 5)

    _title(document, "۵. انتخاب feature و قرارداد preprocessing")
    _body(document, "catalog featureها را با protocol، توضیح و cost ثبت می‌کند. `select analyze` وضعیت usable/constant/near_constant/low_coverage/not_observed/not_implemented را محاسبه می‌کند. کاربر پس از دیدن این evidence یک profile نسخه‌دار ایجاد می‌کند.")
    _diagram(document, diagrams["contract"], "شکل ۳ — profile نسخه‌دار تا model contract inference")
    _body(document, "`preprocessing/pipeline.py` فقط روی training rows fit می‌شود. numericها impute و robust-scale، categoryها محدود و one-hot encode می‌شوند. manifest تمام featureهای انتخابی، حذف‌شده، vocabulary، مدل columnهای transform‌شده و ترتیب آن‌ها را ثبت می‌کند.")
    _table(document, ["artifact", "چرا مهم است؟"], [["profile JSON", "انتخاب feature کاربر و نسخهٔ آن."], ["quality report", "دلیل قابل دفاع برای نگه‌داشتن یا حذف feature."], ["pipeline.joblib", "خود transformation فیت‌شده."], ["manifest JSON", "schema و column order مورد انتظار مدل."], ["model-contract.json", "اتصال profile/preprocessor به مدل deploy شده."]])
    _page_break(document, 6)

    _title(document, "۶. pipeline مدل، ارزیابی و خروجی تحلیلی")
    _body(document, "پس از فریز feature contract، مدل دومرحله‌ای اجرا می‌شود. Stage 1 normal-only LSTM-AE و Isolation Forest را fit می‌کند، threshold را روی normal calibration می‌سازد و در صورت label captureهای کافی گیت supervised سبک را اضافه می‌کند. Stage 2 نوع حمله را فقط برای anomaly candidateها پیش‌بینی می‌کند.")
    _diagram(document, diagrams["two_stage"], "شکل ۴ — pipeline مدل در مرز ماژول")
    _body(document, "هر خروجی تحلیلی بیرون از dashboard نیز وجود دارد: score پکت، evidence feature، history LSTM، readiness و metric RF، prediction end-to-end، latency، threshold sensitivity و resource usage. dashboard همان Parquet/JSONها را visual می‌کند.")
    _table(document, ["خروجی", "کاربرد سیستم دیگر"], [["scores.parquet", "سامانهٔ alert، SIEM یا پروندهٔ incident."], ["feature-evidence.parquet", "توضیح انسانی/گزارش SOC."], ["end-to-end-predictions.parquet", "نتیجهٔ normal/anomaly/type برای API یا DB."], ["model-contract.json", "بارگذاری امن مدل در worker inference."], ["training-summary.json", "نمایش KPI و رصد release مدل."]])
    _page_break(document, 7)

    _title(document, "۷. ذخیره‌سازی artifact و ClickHouse")
    _heading(document, "منبع حقیقت")
    _body(document, "`storage/artifacts.py` ابتدا artifact را به‌صورت فایل محلی می‌نویسد و در `catalog/artifacts.parquet` ثبت می‌کند. این انتخاب باعث می‌شود تغییر یا قطع ClickHouse، مسیر extraction/train/inference را متوقف نکند. schema hash و description به مصرف‌کننده برای اعتبارسنجی کمک می‌کنند.")
    _body(document, "ClickHouseMirror فقط وقتی در config فعال شود ساخته می‌شود. data frameهای artifact را به جدول‌های سازگار mirror می‌کند، اما هیچ pipeline برای خواندن حیاتی به آن وابسته نیست. credential با environment variable خوانده می‌شود، نه از source code.")
    _diagram(document, diagrams["operations"], "شکل ۵ — مصرف artifactها توسط dashboard، CLI، API و ClickHouse")
    _page_break(document, 8)

    _title(document, "۸. داشبورد، API، CLI و مانیتورینگ")
    _body(document, "`dashboard/workbench.py` یک workbench فارسی RTL است. تحلیل کلی و per-protocol، نمودارهای mapping/feature/model، جدول‌های evidence، نمودار pipeline، download artifact، live training monitor و resource monitor در این لایه است. dashboard با انتخاب run و protocol فقط فایل‌های local را می‌خواند.")
    _body(document, "`cli.py` همان قابلیت‌ها را بدون dashboard عرضه می‌کند و logهای طولانی را انگلیسی نگه می‌دارد تا برای terminal/server قابل پردازش باشند. `service.py` façade برنامه‌نویسی است که برای product دیگر، API داخلی یا batch worker مناسب است.")
    _table(document, ["رابط", "ورودی", "خروجی"], [["CLI", "path/config/options", "artifact path، summary و event log."], ["Dashboard", "artifact root و انتخاب اپراتور", "نمودار/جدول/دانلود بدون تغییر source data."], ["AnomalyService", "PCAP/profile/model artifact", "pathهای صریح model/scores/metrics."], ["ClickHouse", "mirror اختیاری frameها", "table برای query سازمانی، نه truth اصلی."]])
    _page_break(document, 9)

    _title(document, "۹. استقرار، عملیات و assurance کیفیت")
    _heading(document, "استقرار Linux/WSL")
    _body(document, "پروژه با Python 3.11/3.12 و `uv` قابل اجرا است. برای server بدون اینترنت، binary `uv`، Python runtime، uv cache/wheelhouse و کل پوشهٔ پروژه از Windows/WSL منتقل می‌شود، سپس `uv sync --offline --frozen` اجرا می‌شود. runtime پیش‌فرض CPU است.")
    _heading(document, "سیاست منابع")
    _body(document, "`runtime.cpu_workers`، `memory_limit_gb`، `capture.streaming_*` و `models.max_source_rows` مرزهای اجرای سنگین را کنترل می‌کنند. extraction تمام داده می‌تواند disk-backed باشد، درحالی‌که model fit با sample نماینده و محدود انجام می‌شود. progress دوام‌دار مانع انتظار بدون اطلاع کاربر می‌شود.")
    _heading(document, "تست و تغییر امن")
    for item in [
        "تست protocol parser، mapping، cache، selection، preprocessing، مدل‌ها و dashboard در `tests/` جدا شده‌اند.",
        "افزودن protocol مستلزم parser، feature catalog، config، تست و راهنمای جدید است.",
        "تغییر schema artifact باید هم‌زمان writer، catalog، dashboard reader و integration guide را به‌روز کند.",
        "قبل از release مدل، mapping audit، normal holdout FPR و capture-held-out attack metrics باید بازبینی شوند.",
    ]:
        _body(document, item, bullet=True)

    document.save(output)
    return output


def main() -> None:
    diagrams = build_diagrams()
    model_report = build_model_report(diagrams)
    module_report = build_module_report(diagrams)
    print(f"Created: {model_report}")
    print(f"Created: {module_report}")


if __name__ == "__main__":
    main()
